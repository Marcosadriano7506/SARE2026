from io import BytesIO

from openpyxl import load_workbook

from app.extensions import db
from app.models import (
    Answer,
    ApplicationStatus,
    ClassApplication,
    ClassRoom,
    Evaluation,
    Question,
    School,
    Skill,
    Student,
    StudentPresence,
    StudentRecord,
    SubjectArea,
    Test,
    User,
    UserRole,
)
from app.services.results_workbook import build_results_workbook


def test_results_workbook_has_grade_tabs_and_per_student_skill_scores(app):
    with app.app_context():
        applicator = User(
            name="Aplicador",
            username="export-app",
            role=UserRole.APPLICATOR,
        )
        applicator.set_password("secret123")

        evaluation = Evaluation(name="SARE EXPORT", school_year=2026)
        school = School(name="Escola Exportação")
        classroom = ClassRoom(
            school=school,
            evaluation=evaluation,
            grade=5,
            name="5º Ano A",
            access_code="EXP5A",
        )

        present = Student(
            classroom=classroom,
            name="Aluno Presente",
            external_id="P1",
        )
        absent = Student(
            classroom=classroom,
            name="Aluno Ausente",
            external_id="A1",
        )

        lp_skill = Skill(
            code="D01",
            grade=5,
            subject=SubjectArea.PORTUGUESE,
        )
        math_skill = Skill(
            code="D02",
            grade=5,
            subject=SubjectArea.MATHEMATICS,
        )
        lp_test = Test(
            evaluation=evaluation,
            grade=5,
            subject=SubjectArea.PORTUGUESE,
            title="LP 5º",
        )
        math_test = Test(
            evaluation=evaluation,
            grade=5,
            subject=SubjectArea.MATHEMATICS,
            title="MAT 5º",
        )
        lp_question = Question(
            test=lp_test,
            number=1,
            skill=lp_skill,
            correct_option="A",
        )
        math_question = Question(
            test=math_test,
            number=1,
            skill=math_skill,
            correct_option="B",
        )

        application = ClassApplication(
            classroom=classroom,
            applicator=applicator,
            status=ApplicationStatus.FINALIZED,
        )

        db.session.add_all(
            [
                applicator,
                evaluation,
                school,
                classroom,
                present,
                absent,
                lp_skill,
                math_skill,
                lp_test,
                math_test,
                lp_question,
                math_question,
                application,
            ]
        )
        db.session.flush()

        present_record = StudentRecord(
            class_application=application,
            student=present,
            presence=StudentPresence.PRESENT,
            saved_by=applicator.id,
        )
        present_record.answers = [
            Answer(question=lp_question, selected_option="A"),
            Answer(question=math_question, selected_option="C"),
        ]
        absent_record = StudentRecord(
            class_application=application,
            student=absent,
            presence=StudentPresence.ABSENT,
            saved_by=applicator.id,
        )
        db.session.add_all([present_record, absent_record])
        db.session.commit()

        data = build_results_workbook(evaluation)

    workbook = load_workbook(BytesIO(data), data_only=True)
    assert workbook.sheetnames[:8] == [
        "2º ANO",
        "3º ANO",
        "4º ANO",
        "5º ANO",
        "6º ANO",
        "7º ANO",
        "8º ANO",
        "9º ANO",
    ]

    sheet = workbook["5º ANO"]
    headers = [cell.value for cell in sheet[1]]
    assert "D01" in headers
    assert "D02" in headers
    assert "TOTAL DE ACERTOS LÍNGUA PORTUGUESA" in headers
    assert "TOTAL DE ACERTOS MATEMÁTICA" in headers

    rows_by_name = {
        sheet.cell(row=row, column=5).value: row
        for row in range(2, sheet.max_row + 1)
    }

    present_row = rows_by_name["Aluno Presente"]
    absent_row = rows_by_name["Aluno Ausente"]

    d01_col = headers.index("D01") + 1
    d02_col = headers.index("D02") + 1
    lp_total_col = headers.index("TOTAL DE ACERTOS LÍNGUA PORTUGUESA") + 1
    math_total_col = headers.index("TOTAL DE ACERTOS MATEMÁTICA") + 1
    overall_total_col = headers.index("TOTAL DE ACERTOS") + 1

    assert sheet.cell(present_row, 4).value == "SIM"
    assert sheet.cell(present_row, d01_col).value == 1
    assert sheet.cell(present_row, d02_col).value == 0
    assert sheet.cell(present_row, lp_total_col).value == 1
    assert sheet.cell(present_row, math_total_col).value == 0
    assert sheet.cell(present_row, overall_total_col).value == 1

    assert sheet.cell(absent_row, 4).value == "NÃO"
    assert sheet.cell(absent_row, d01_col).value is None
    assert sheet.cell(absent_row, d02_col).value is None


def test_duplicate_skill_codes_are_disambiguated_with_question_number(app):
    with app.app_context():
        evaluation = Evaluation(name="SARE DUP", school_year=2026)
        skill = Skill(
            code="D01",
            grade=5,
            subject=SubjectArea.PORTUGUESE,
        )
        test = Test(
            evaluation=evaluation,
            grade=5,
            subject=SubjectArea.PORTUGUESE,
            title="LP",
        )
        test.questions = [
            Question(number=1, skill=skill, correct_option="A"),
            Question(number=2, skill=skill, correct_option="B"),
        ]
        db.session.add_all([evaluation, skill, test])
        db.session.commit()

        data = build_results_workbook(evaluation)

    workbook = load_workbook(BytesIO(data), data_only=True)
    headers = [cell.value for cell in workbook["5º ANO"][1]]
    assert "D01 (Q1)" in headers
    assert "D01 (Q2)" in headers
