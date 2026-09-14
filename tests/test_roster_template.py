from io import BytesIO

from openpyxl import load_workbook

from app.extensions import db
from app.models import User, UserRole


def _admin(app):
    with app.app_context():
        user = User(
            name="Admin Modelo",
            username="admin-modelo",
            role=UserRole.ADMIN,
        )
        user.set_password("secret123")
        db.session.add(user)
        db.session.commit()
        return user.id


def _login(client, user_id):
    with client.session_transaction() as session:
        session["_user_id"] = str(user_id)
        session["_fresh"] = True


def test_download_official_roster_template(app, client):
    admin_id = _admin(app)
    _login(client, admin_id)

    response = client.get("/coordenacao/base/modelo.xlsx")

    assert response.status_code == 200
    assert response.mimetype == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    workbook = load_workbook(BytesIO(response.data))
    assert workbook.sheetnames == ["BASE_SARE", "INSTRUCOES"]

    sheet = workbook["BASE_SARE"]
    assert [cell.value for cell in sheet[1]] == [
        "ESCOLA",
        "ANO",
        "TURMA",
        "ALUNO",
        "MATRÍCULA",
    ]
    assert sheet["E2"].value == "000001"
    assert sheet["E2"].number_format == "@"

    instructions = workbook["INSTRUCOES"]
    assert "MODELO OFICIAL" in instructions["A1"].value
