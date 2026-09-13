from __future__ import annotations

import secrets
import unicodedata
from dataclasses import dataclass
from io import BytesIO

from openpyxl import load_workbook

from app.extensions import db
from app.models import ClassRoom, Evaluation, School, Student


HEADER_ALIASES = {
    "school": {"ESCOLA", "SCHOOL", "UNIDADE ESCOLAR"},
    "grade": {"ANO", "GRADE", "SERIE", "SÉRIE", "ANO/SERIE", "ANO/SÉRIE"},
    "class_name": {"TURMA", "CLASS", "CLASSE"},
    "student": {"ALUNO", "ESTUDANTE", "STUDENT", "NOME DO ALUNO", "NOME DO ESTUDANTE"},
    "external_id": {"MATRICULA", "MATRÍCULA", "ID", "CODIGO", "CÓDIGO", "ID ALUNO"},
}


@dataclass(frozen=True)
class ParsedRosterRow:
    row_number: int
    school: str
    grade: int
    class_name: str
    student: str
    external_id: str | None


@dataclass(frozen=True)
class ImportResult:
    schools_created: int
    classes_created: int
    students_created: int
    rows_processed: int


class RosterImportError(ValueError):
    def __init__(self, errors: list[str]):
        self.errors = errors
        super().__init__("; ".join(errors))


def _normalize(value) -> str:
    text = "" if value is None else str(value).strip()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return " ".join(text.upper().split())


def _header_map(values: list) -> dict[str, int]:
    normalized = [_normalize(value) for value in values]
    mapping: dict[str, int] = {}
    for key, aliases in HEADER_ALIASES.items():
        normalized_aliases = {_normalize(alias) for alias in aliases}
        for idx, value in enumerate(normalized):
            if value in normalized_aliases:
                mapping[key] = idx
                break
    return mapping


def _grade(value) -> int:
    normalized = _normalize(value)
    digits = "".join(ch for ch in normalized if ch.isdigit())
    if not digits:
        raise ValueError("ano inválido")
    grade = int(digits)
    if grade < 2 or grade > 9:
        raise ValueError("ano deve estar entre 2 e 9")
    return grade


def parse_roster_xlsx(content: bytes) -> list[ParsedRosterRow]:
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise RosterImportError(["Arquivo Excel inválido ou corrompido."]) from exc

    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    try:
        header = list(next(rows))
    except StopIteration as exc:
        raise RosterImportError(["A planilha está vazia."]) from exc

    mapping = _header_map(header)
    missing = [key for key in ("school", "grade", "class_name", "student") if key not in mapping]
    if missing:
        friendly = {
            "school": "ESCOLA",
            "grade": "ANO",
            "class_name": "TURMA",
            "student": "ALUNO",
        }
        raise RosterImportError([
            "Colunas obrigatórias ausentes: " + ", ".join(friendly[item] for item in missing)
        ])

    parsed: list[ParsedRosterRow] = []
    errors: list[str] = []

    for row_number, row in enumerate(rows, start=2):
        values = list(row)
        if not any(value not in (None, "") for value in values):
            continue

        def cell(key):
            idx = mapping.get(key)
            return values[idx] if idx is not None and idx < len(values) else None

        school = str(cell("school") or "").strip()
        class_name = str(cell("class_name") or "").strip()
        student = str(cell("student") or "").strip()
        external_raw = cell("external_id")
        external_id = str(external_raw).strip() if external_raw not in (None, "") else None

        row_errors = []
        if not school:
            row_errors.append("escola vazia")
        if not class_name:
            row_errors.append("turma vazia")
        if not student:
            row_errors.append("aluno vazio")
        try:
            grade = _grade(cell("grade"))
        except ValueError as exc:
            grade = 0
            row_errors.append(str(exc))

        if row_errors:
            errors.append(f"Linha {row_number}: " + ", ".join(row_errors))
            continue

        parsed.append(
            ParsedRosterRow(
                row_number=row_number,
                school=school,
                grade=grade,
                class_name=class_name,
                student=student,
                external_id=external_id,
            )
        )

    if errors:
        raise RosterImportError(errors[:30])
    if not parsed:
        raise RosterImportError(["Nenhum estudante válido encontrado na planilha."])
    return parsed


def _new_class_code() -> str:
    return secrets.token_hex(4).upper()


def import_roster(evaluation: Evaluation, rows: list[ParsedRosterRow]) -> ImportResult:
    schools_created = 0
    classes_created = 0
    students_created = 0

    school_cache = {school.name.upper(): school for school in School.query.all()}
    class_cache: dict[tuple[int, int, str], ClassRoom] = {}

    for classroom in ClassRoom.query.filter_by(evaluation_id=evaluation.id).all():
        class_cache[(classroom.school_id, classroom.grade, classroom.name.upper())] = classroom

    for row in rows:
        school_key = row.school.upper()
        school = school_cache.get(school_key)
        if school is None:
            school = School(name=row.school)
            db.session.add(school)
            db.session.flush()
            school_cache[school_key] = school
            schools_created += 1

        class_key = (school.id, row.grade, row.class_name.upper())
        classroom = class_cache.get(class_key)
        if classroom is None:
            classroom = ClassRoom(
                school=school,
                evaluation=evaluation,
                grade=row.grade,
                name=row.class_name,
                access_code=_new_class_code(),
            )
            db.session.add(classroom)
            db.session.flush()
            class_cache[class_key] = classroom
            classes_created += 1

        if row.external_id:
            existing = Student.query.filter_by(
                class_id=classroom.id, external_id=row.external_id
            ).first()
        else:
            existing = Student.query.filter(
                Student.class_id == classroom.id,
                db.func.upper(Student.name) == row.student.upper(),
            ).first()

        if existing is None:
            db.session.add(
                Student(
                    classroom=classroom,
                    external_id=row.external_id,
                    name=row.student,
                )
            )
            students_created += 1

    db.session.flush()
    return ImportResult(
        schools_created=schools_created,
        classes_created=classes_created,
        students_created=students_created,
        rows_processed=len(rows),
    )
