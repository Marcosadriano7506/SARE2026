from __future__ import annotations

import unicodedata
from dataclasses import dataclass
from io import BytesIO

from openpyxl import load_workbook

from app.extensions import db
from app.models import Evaluation, Question, Skill, SubjectArea, Test


HEADER_ALIASES = {
    "grade": {"ANO", "SERIE", "SÉRIE", "GRADE"},
    "subject": {"COMPONENTE", "DISCIPLINA", "SUBJECT"},
    "question": {"QUESTAO", "QUESTÃO", "ITEM", "NUMERO", "NÚMERO"},
    "skill": {"HABILIDADE", "DESCRITOR", "SKILL"},
    "answer": {"GABARITO", "RESPOSTA", "ANSWER"},
    "skill_description": {
        "DESCRICAO DA HABILIDADE",
        "DESCRIÇÃO DA HABILIDADE",
        "DESCRICAO",
        "DESCRIÇÃO",
    },
    "expected_outcome": {
        "O QUE SE ESPERA",
        "EXPECTATIVA",
        "RESULTADO ESPERADO",
        "EXPECTATIVA DE APRENDIZAGEM",
        "EXPECTED OUTCOME",
    },
}


@dataclass(frozen=True)
class ParsedKeyRow:
    row_number: int
    grade: int
    subject: SubjectArea
    question_number: int
    skill_code: str
    correct_option: str
    skill_description: str | None
    expected_outcome: str | None


@dataclass(frozen=True)
class KeyImportResult:
    tests_created: int
    skills_created: int
    questions_created: int
    questions_updated: int
    rows_processed: int


class AnswerKeyImportError(ValueError):
    def __init__(self, errors: list[str]):
        self.errors = errors
        super().__init__("; ".join(errors))


def _normalize(value) -> str:
    text = "" if value is None else str(value).strip()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return " ".join(text.upper().split())


def _header_map(values: list) -> dict[str, int]:
    normalized = [_normalize(value) for value in values]
    mapping: dict[str, int] = {}
    for key, aliases in HEADER_ALIASES.items():
        normalized_aliases = {_normalize(alias) for alias in aliases}
        for idx, value in enumerate(normalized):
            if value in normalized_aliases:
                mapping[key] = idx
                break
    return mapping


def _grade(value) -> int:
    digits = "".join(ch for ch in _normalize(value) if ch.isdigit())
    if not digits:
        raise ValueError("ano inválido")
    grade = int(digits)
    if grade < 2 or grade > 9:
        raise ValueError("ano deve estar entre 2 e 9")
    return grade


def _subject(value) -> SubjectArea:
    normalized = _normalize(value)
    if normalized in {"LP", "PORTUGUES", "LINGUA PORTUGUESA", "PORTUGUESE"}:
        return SubjectArea.PORTUGUESE
    if normalized in {"MA", "MAT", "MATEMATICA", "MATHEMATICS"}:
        return SubjectArea.MATHEMATICS
    raise ValueError("componente deve ser LP ou MATEMÁTICA")


def parse_answer_key_xlsx(content: bytes) -> list[ParsedKeyRow]:
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise AnswerKeyImportError(["Arquivo Excel inválido ou corrompido."]) from exc

    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    try:
        header = list(next(rows))
    except StopIteration as exc:
        raise AnswerKeyImportError(["A planilha está vazia."]) from exc

    mapping = _header_map(header)
    required = ("grade", "subject", "question", "skill", "answer")
    missing = [key for key in required if key not in mapping]
    if missing:
        labels = {
            "grade": "ANO",
            "subject": "COMPONENTE",
            "question": "QUESTÃO",
            "skill": "HABILIDADE",
            "answer": "GABARITO",
        }
        raise AnswerKeyImportError([
            "Colunas obrigatórias ausentes: " + ", ".join(labels[key] for key in missing)
        ])

    parsed: list[ParsedKeyRow] = []
    errors: list[str] = []
    seen: set[tuple[int, SubjectArea, int]] = set()

    for row_number, row in enumerate(rows, start=2):
        values = list(row)
        if not any(value not in (None, "") for value in values):
            continue

        def cell(key):
            idx = mapping.get(key)
            return values[idx] if idx is not None and idx < len(values) else None

        row_errors = []
        try:
            grade = _grade(cell("grade"))
        except ValueError as exc:
            grade = 0
            row_errors.append(str(exc))

        try:
            subject = _subject(cell("subject"))
        except ValueError as exc:
            subject = SubjectArea.PORTUGUESE
            row_errors.append(str(exc))

        try:
            question_number = int(cell("question"))
            if question_number <= 0:
                raise ValueError
        except (TypeError, ValueError):
            question_number = 0
            row_errors.append("número da questão inválido")

        skill_code = str(cell("skill") or "").strip().upper()
        if not skill_code:
            row_errors.append("habilidade vazia")

        correct_option = _normalize(cell("answer"))
        if correct_option not in {"A", "B", "C", "D"}:
            row_errors.append("gabarito deve ser A, B, C ou D")

        description_raw = cell("skill_description")
        description = (
            str(description_raw).strip()
            if description_raw not in (None, "")
            else None
        )

        expected_outcome_raw = cell("expected_outcome")
        expected_outcome = (
            str(expected_outcome_raw).strip()
            if expected_outcome_raw not in (None, "")
            else None
        )

        identity = (grade, subject, question_number)
        if not row_errors and identity in seen:
            row_errors.append("questão duplicada para o mesmo ano/componente")

        if row_errors:
            errors.append(f"Linha {row_number}: " + ", ".join(row_errors))
            continue

        seen.add(identity)
        parsed.append(
            ParsedKeyRow(
                row_number=row_number,
                grade=grade,
                subject=subject,
                question_number=question_number,
                skill_code=skill_code,
                correct_option=correct_option,
                skill_description=description,
                expected_outcome=expected_outcome,
            )
        )

    if errors:
        raise AnswerKeyImportError(errors[:30])
    if not parsed:
        raise AnswerKeyImportError(["Nenhuma questão válida encontrada na planilha."])
    return parsed


def import_answer_key(
    evaluation: Evaluation, rows: list[ParsedKeyRow]
) -> KeyImportResult:
    tests_created = 0
    skills_created = 0
    questions_created = 0
    questions_updated = 0

    test_cache: dict[tuple[int, SubjectArea], Test] = {}
    skill_cache: dict[tuple[int, SubjectArea, str], Skill] = {}

    for row in rows:
        test_key = (row.grade, row.subject)
        test = test_cache.get(test_key)
        if test is None:
            test = Test.query.filter_by(
                evaluation_id=evaluation.id,
                grade=row.grade,
                subject=row.subject,
            ).first()
            if test is None:
                subject_label = (
                    "Língua Portuguesa"
                    if row.subject == SubjectArea.PORTUGUESE
                    else "Matemática"
                )
                test = Test(
                    evaluation=evaluation,
                    grade=row.grade,
                    subject=row.subject,
                    title=f"{subject_label} — {row.grade}º Ano",
                )
                db.session.add(test)
                db.session.flush()
                tests_created += 1
            test_cache[test_key] = test

        skill_key = (row.grade, row.subject, row.skill_code)
        skill = skill_cache.get(skill_key)
        if skill is None:
            skill = Skill.query.filter_by(
                grade=row.grade,
                subject=row.subject,
                code=row.skill_code,
            ).first()
            if skill is None:
                skill = Skill(
                    code=row.skill_code,
                    grade=row.grade,
                    subject=row.subject,
                    description=row.skill_description,
                    expected_outcome=row.expected_outcome,
                )
                db.session.add(skill)
                db.session.flush()
                skills_created += 1
            else:
                if row.skill_description:
                    skill.description = row.skill_description
                if row.expected_outcome:
                    skill.expected_outcome = row.expected_outcome
            skill_cache[skill_key] = skill

        question = Question.query.filter_by(
            test_id=test.id, number=row.question_number
        ).first()
        if question is None:
            question = Question(
                test=test,
                number=row.question_number,
                skill=skill,
                correct_option=row.correct_option,
            )
            db.session.add(question)
            questions_created += 1
        else:
            question.skill = skill
            question.correct_option = row.correct_option
            questions_updated += 1

    db.session.flush()
    return KeyImportResult(
        tests_created=tests_created,
        skills_created=skills_created,
        questions_created=questions_created,
        questions_updated=questions_updated,
        rows_processed=len(rows),
    )
