from __future__ import annotations

from collections import Counter
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.models import (
    ApplicationStatus,
    StudentPresence,
    SubjectArea,
)
from app.services.analytics import calculate_evaluation_analytics
from app.services.scoring import proficiency_level


THIN = Side(style="thin", color="000000")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

IDENTITY_FILL = PatternFill("solid", fgColor="00B0F0")
LP_FILL = PatternFill("solid", fgColor="D9EAF7")
MATH_FILL = PatternFill("solid", fgColor="FFF2CC")
TOTAL_FILL = PatternFill("solid", fgColor="A8D08D")
PERCENT_FILL = PatternFill("solid", fgColor="FFC000")
GENERAL_FILL = PatternFill("solid", fgColor="D9EAD3")


DECLARATION_LABELS = {
    "PRETO": "PRETO",
    "PARDO": "PARDO",
    "AMARELO": "AMARELO",
    "INDIGENA": "INDÍGENA",
    "BRANCO": "BRANCO",
}


def _skill_headers(questions):
    counts = Counter(
        (question.skill.code if question.skill else f"Q{question.number}")
        for question in questions
    )
    headers = []
    for question in questions:
        base = question.skill.code if question.skill else f"Q{question.number}"
        headers.append(
            f"{base} (Q{question.number})" if counts[base] > 1 else base
        )
    return headers


def _student_record_for(classroom, student):
    application = classroom.application
    if (
        application is None
        or application.status != ApplicationStatus.FINALIZED
    ):
        return None, "PENDENTE"

    record = next(
        (item for item in application.records if item.student_id == student.id),
        None,
    )
    if record is None:
        return None, "PENDENTE"
    if record.presence == StudentPresence.ABSENT:
        return record, "NÃO"
    return record, "SIM"


def _proficiency_from_fraction(value: float | None) -> str | None:
    if value is None:
        return None
    return proficiency_level(value * 100).replace("INTERMEDIARIO", "INTERMEDIÁRIO").replace(
        "AVANCADO", "AVANÇADO"
    )


def _question_score(record, question) -> int:
    answers = {answer.question_id: answer.selected_option for answer in record.answers}
    selected = answers.get(question.id)
    return 1 if selected is not None and selected == question.correct_option else 0


def _write_grade_sheet(workbook: Workbook, evaluation, grade: int):
    title = f"{grade}º ANO"
    sheet = workbook.create_sheet(title)

    tests = sorted(
        [test for test in evaluation.tests if test.grade == grade],
        key=lambda item: item.subject.value,
    )
    lp_questions = []
    math_questions = []
    for test in tests:
        ordered = sorted(test.questions, key=lambda item: item.number)
        if test.subject == SubjectArea.PORTUGUESE:
            lp_questions.extend(ordered)
        elif test.subject == SubjectArea.MATHEMATICS:
            math_questions.extend(ordered)

    headers = [
        "ESCOLA",
        "TURMA",
        "AUTODECLARAÇÃO",
        "PRESENTE",
        "NOME DO ESTUDANTE",
        *_skill_headers(lp_questions),
        "TOTAL DE ACERTOS LÍNGUA PORTUGUESA",
        "PERCENTUAL DE ACERTOS LÍNGUA PORTUGUESA",
        "PROFICIÊNCIA LÍNGUA PORTUGUESA",
        *_skill_headers(math_questions),
        "TOTAL DE ACERTOS MATEMÁTICA",
        "PERCENTUAL DE ACERTOS MATEMÁTICA",
        "PROFICIÊNCIA MATEMÁTICA",
        "TOTAL DE ACERTOS",
        "PERCENTUAL DE ACERTOS",
    ]
    sheet.append(headers)

    lp_start = 6
    lp_end = lp_start + len(lp_questions) - 1
    lp_total_col = lp_start + len(lp_questions)
    lp_percent_col = lp_total_col + 1
    lp_prof_col = lp_total_col + 2

    math_start = lp_prof_col + 1
    math_end = math_start + len(math_questions) - 1
    math_total_col = math_start + len(math_questions)
    math_percent_col = math_total_col + 1
    math_prof_col = math_total_col + 2

    overall_total_col = math_prof_col + 1
    overall_percent_col = overall_total_col + 1

    classrooms = sorted(
        [classroom for classroom in evaluation.classes if classroom.grade == grade],
        key=lambda item: (item.school.name, item.name),
    )

    for classroom in classrooms:
        for student in sorted(classroom.students, key=lambda item: item.name):
            record, presence_label = _student_record_for(classroom, student)

            declaration = None
            lp_scores = [None] * len(lp_questions)
            math_scores = [None] * len(math_questions)
            lp_total = lp_percent = lp_prof = None
            math_total = math_percent = math_prof = None
            overall_total = overall_percent = None

            if record is not None and record.self_declaration is not None:
                declaration = DECLARATION_LABELS.get(
                    record.self_declaration.value,
                    record.self_declaration.value,
                )

            if record is not None and record.presence == StudentPresence.PRESENT:
                lp_scores = [_question_score(record, question) for question in lp_questions]
                math_scores = [
                    _question_score(record, question) for question in math_questions
                ]

                lp_total = sum(lp_scores)
                math_total = sum(math_scores)
                lp_percent = (
                    lp_total / len(lp_questions) if lp_questions else None
                )
                math_percent = (
                    math_total / len(math_questions) if math_questions else None
                )
                lp_prof = _proficiency_from_fraction(lp_percent)
                math_prof = _proficiency_from_fraction(math_percent)

                overall_total = lp_total + math_total
                total_items = len(lp_questions) + len(math_questions)
                overall_percent = (
                    overall_total / total_items if total_items else None
                )

            sheet.append(
                [
                    classroom.school.name,
                    classroom.name,
                    declaration,
                    presence_label,
                    student.name,
                    *lp_scores,
                    lp_total,
                    lp_percent,
                    lp_prof,
                    *math_scores,
                    math_total,
                    math_percent,
                    math_prof,
                    overall_total,
                    overall_percent,
                ]
            )

    # Header styling inspired by the reference workbook supplied by the user.
    for col in range(1, 6):
        cell = sheet.cell(row=1, column=col)
        cell.fill = IDENTITY_FILL
        cell.font = Font(bold=True, name="Arial", size=11)
        cell.border = BORDER

    if lp_questions:
        for col in range(lp_start, lp_end + 1):
            sheet.cell(row=1, column=col).fill = LP_FILL
    if math_questions:
        for col in range(math_start, math_end + 1):
            sheet.cell(row=1, column=col).fill = MATH_FILL

    for col in [lp_total_col, math_total_col, overall_total_col]:
        sheet.cell(row=1, column=col).fill = TOTAL_FILL
    for col in [
        lp_percent_col,
        lp_prof_col,
        math_percent_col,
        math_prof_col,
        overall_percent_col,
    ]:
        sheet.cell(row=1, column=col).fill = PERCENT_FILL

    for cell in sheet[1]:
        cell.font = Font(bold=True, name="Arial", size=10)
        cell.border = BORDER
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.border = BORDER
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(
                horizontal="center" if cell.column >= 3 else "left",
                vertical="center",
                wrap_text=False,
            )

    for row_idx in range(2, sheet.max_row + 1):
        for col in [lp_percent_col, math_percent_col, overall_percent_col]:
            sheet.cell(row=row_idx, column=col).number_format = "0.00%"

    # Keep the identity columns visible when scrolling through many skills.
    sheet.freeze_panes = "F2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 42

    widths = {
        1: 36,
        2: 22,
        3: 20,
        4: 12,
        5: 34,
    }
    for col_idx, width in widths.items():
        sheet.column_dimensions[get_column_letter(col_idx)].width = width

    for col_idx in range(6, overall_percent_col + 1):
        header = str(sheet.cell(row=1, column=col_idx).value or "")
        if "PERCENTUAL" in header or "PROFICIÊNCIA" in header or "TOTAL DE ACERTOS" in header:
            sheet.column_dimensions[get_column_letter(col_idx)].width = 24
        else:
            sheet.column_dimensions[get_column_letter(col_idx)].width = 16

    return sheet


def _write_analytical_sheets(workbook: Workbook, evaluation, analytics):
    summary = workbook.create_sheet("RESUMO")
    summary.append(["AVALIAÇÃO", evaluation.name])
    summary.append(["TURMAS FINALIZADAS", analytics.finalized_classes])
    summary.append(["PRESENTES", analytics.present_students])
    summary.append(["AUSENTES", analytics.absent_students])
    summary.append(["RESULTADO GERAL (%)", analytics.total_percent])
    summary.append(["LÍNGUA PORTUGUESA (%)", analytics.lp_percent])
    summary.append(["MATEMÁTICA (%)", analytics.math_percent])

    students_sheet = workbook.create_sheet("ALUNOS")
    students_sheet.append([
        "ANO",
        "ESCOLA",
        "TURMA",
        "ALUNO",
        "GERAL (%)",
        "LP (%)",
        "MATEMÁTICA (%)",
        "NÍVEL",
    ])
    for item in analytics.students:
        students_sheet.append([
            item.grade,
            item.school_name,
            item.class_name,
            item.student_name,
            item.total_percent,
            item.lp_percent,
            item.math_percent,
            item.level,
        ])

    skills_sheet = workbook.create_sheet("HABILIDADES")
    skills_sheet.append([
        "ANO",
        "COMPONENTE",
        "HABILIDADE",
        "ACERTOS",
        "OPORTUNIDADES",
        "%",
    ])
    for item in analytics.skills:
        skills_sheet.append([
            item.grade,
            "LP" if item.subject == "PORTUGUESE" else "MATEMÁTICA",
            item.code,
            item.correct,
            item.opportunities,
            item.percent,
        ])

    questions_sheet = workbook.create_sheet("QUESTOES")
    questions_sheet.append([
        "ANO",
        "COMPONENTE",
        "QUESTÃO",
        "HABILIDADE",
        "ACERTOS",
        "OPORTUNIDADES",
        "%",
    ])
    for item in analytics.questions:
        questions_sheet.append([
            item.grade,
            "LP" if item.subject == "PORTUGUESE" else "MATEMÁTICA",
            item.number,
            item.skill_code,
            item.correct,
            item.opportunities,
            item.percent,
        ])

    schools_sheet = workbook.create_sheet("ESCOLAS")
    schools_sheet.append([
        "POSIÇÃO",
        "ESCOLA",
        "PRESENTES",
        "ACERTOS",
        "ITENS",
        "%",
    ])
    for position, item in enumerate(analytics.schools, start=1):
        schools_sheet.append([
            position,
            item.name,
            item.present_students,
            item.correct,
            item.items,
            item.percent,
        ])

    classes_sheet = workbook.create_sheet("TURMAS")
    classes_sheet.append([
        "POSIÇÃO",
        "TURMA",
        "PRESENTES",
        "ACERTOS",
        "ITENS",
        "%",
    ])
    for position, item in enumerate(analytics.classes, start=1):
        classes_sheet.append([
            position,
            item.name,
            item.present_students,
            item.correct,
            item.items,
            item.percent,
        ])

    declarations_sheet = workbook.create_sheet("AUTODECLARACAO")
    declarations_sheet.append([
        "ESCOPO",
        "LOCAL",
        "ANO",
        "AUTODECLARAÇÃO",
        "ESTUDANTES",
        "ACERTOS",
        "ITENS",
        "%",
    ])
    for item in analytics.declarations:
        declarations_sheet.append([
            item.scope_type,
            item.scope_name,
            item.grade,
            item.declaration,
            item.students,
            item.correct,
            item.items,
            item.percent,
        ])

    absences_sheet = workbook.create_sheet("AUSENTES")
    absences_sheet.append(["ANO", "ESCOLA", "TURMA", "ALUNO"])
    for item in analytics.absences:
        absences_sheet.append([
            item.grade,
            item.school_name,
            item.class_name,
            item.student_name,
        ])

    for sheet in [
        summary,
        students_sheet,
        skills_sheet,
        questions_sheet,
        schools_sheet,
        classes_sheet,
        declarations_sheet,
        absences_sheet,
    ]:
        sheet.freeze_panes = "A2"
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = GENERAL_FILL
            cell.border = BORDER
        for column_cells in sheet.columns:
            max_length = max(
                len(str(cell.value)) if cell.value is not None else 0
                for cell in column_cells
            )
            sheet.column_dimensions[column_cells[0].column_letter].width = min(
                max_length + 2, 45
            )


def build_results_workbook(evaluation) -> bytes:
    analytics = calculate_evaluation_analytics(evaluation)
    workbook = Workbook()

    # Remove the automatic blank sheet. Detailed grade sheets come first,
    # matching the structure of the supplied reference workbook.
    workbook.remove(workbook.active)

    for grade in range(2, 10):
        _write_grade_sheet(workbook, evaluation, grade)

    _write_analytical_sheets(workbook, evaluation, analytics)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output.getvalue()
